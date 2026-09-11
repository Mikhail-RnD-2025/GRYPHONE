# -*- coding: utf-8 -*-
"""
app/services/camera_import_service.py
======================================
Сервис импорта/экспорта камер из Excel и JSON.
PATCH-190: перенос логики из import_from_excel.py внутрь проекта.
"""
import json
import sqlite3
from pathlib import Path
from typing import Dict, List, Any, Optional
from urllib.parse import urlparse
import logging

logger = logging.getLogger(__name__)


# ============================================================================
# Константы и маппинг
# ============================================================================

CAMERA_FIELDS = {
    'id':         {'type': 'str',  'required': True,  'default': None},
    'name':       {'type': 'str',  'required': False, 'default': ''},
    'login':      {'type': 'str',  'required': False, 'default': ''},
    'pass':       {'type': 'str',  'required': False, 'default': ''},
    'ipaddress':  {'type': 'str',  'required': True,  'default': None},
    'port':       {'type': 'str',  'required': False, 'default': '554'},
    'main_url':   {'type': 'str',  'required': True,  'default': None},
    'sub_url':    {'type': 'str',  'required': False, 'default': ''},
    'sub2_url':   {'type': 'str',  'required': False, 'default': ''},
    'enabled':    {'type': 'bool', 'required': False, 'default': True},
    'comment':    {'type': 'str',  'required': False, 'default': ''},
    'audio':      {'type': 'bool', 'required': False, 'default': True},
    'location':   {'type': 'str',  'required': False, 'default': ''},
}

# Маппинг возможных названий колонок (нижний регистр)
COLUMN_MAPPING = {
    'id': [
        'id', 'ID', 'идентификатор', 'id камеры', 'camera id',
        'номер', 'код', 'code',
    ],
    'name': [
        'name', 'имя', 'название', 'камера', 'camera name',
        'название камеры', 'имя камеры',
    ],
    'login': [
        'login', 'логин', 'пользователь', 'user', 'username',
        'имя пользователя',
    ],
    'pass': [
        'pass', 'password', 'пароль', 'pwd',
    ],
    'ipaddress': [
        'ipaddress', 'ip', 'ip адрес', 'ip-адрес', 'адрес', 'address',
        'ip address', 'ip_address',
    ],
    'port': [
        'port', 'порт',
    ],
    'main_url': [
        'main_url', 'main url', 'основной_url', 'основной url',
        'url', 'ссылка', 'основная ссылка', 'поток', 'stream',
        'rtsp', 'main', 'поток основной', 'основной поток',
        'main_path', 'main path', 'путь основного', 'путь основного потока',
    ],
    'sub_url': [
        'sub_url', 'sub url', 'дополнительный_url', 'дополнительный url',
        'sub', 'суб', 'субпоток', 'substream', 'sub stream',
        'поток дополнительный', 'дополнительный поток',
        'sub_path', 'sub path', 'путь суб', 'путь субпотока',
    ],
    'sub2_url': [
        'sub2_url', 'sub2 url', 'sub2', 'суб2', 'sub2_path', 'sub2 path',
        'путь sub2', 'путь sub2 потока',  # PATCH-196: заголовок экспорта
    ],
    'enabled': [
        'enabled', 'включена', 'активна', 'статус', 'status', 'активно',
        'включено', 'active',
    ],
    'comment': [
        'comment', 'комментарий', 'описание', 'description', 'desc',
        'примечание', 'note',
    ],
    'audio': [
        'audio', 'аудио', 'звук', 'sound', 'audio_enabled',
    ],
    'location': [
        'location', 'местоположение', 'расположение', 'локация', 'location_name',
        'где', 'where', 'место',
    ],
}


# ============================================================================
# Вспомогательные функции
# ============================================================================

def clean_str(value: Any) -> str:
    """Очищает строковое значение от лишних пробелов."""
    if value is None:
        return ''
    return str(value).strip()


def clean_name(value: Any) -> str:
    """Очищает имя камеры от лишних пробелов и дефиса в начале."""
    cleaned = clean_str(value)
    if cleaned.startswith('-'):
        cleaned = cleaned[1:].strip()
    return cleaned


def clean_url(value: Any) -> str:
    """Очищает RTSP-ссылку от лишних пробелов (включая внутри)."""
    if value is None:
        return ''
    return str(value).strip().replace(' ', '')


def parse_bool(value: Any, default: bool = True) -> bool:
    """Парсит булево значение из разных форматов."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)

    value_str = clean_str(value).lower()
    if not value_str:
        return default

    true_values = ['true', '1', 'да', 'yes', 'вкл', 'включена', 'активна', 'включено']
    false_values = ['false', '0', 'нет', 'no', 'выкл', 'выключена', 'неактивна', 'выключено']

    if value_str in true_values:
        return True
    if value_str in false_values:
        return False
    return default


def parse_rtsp_url(url: str) -> Optional[Dict[str, str]]:
    """
    Парсит полный RTSP URL на части.
    Возвращает dict с login, pass, ipaddress, port, path или None.
    """
    if not url or not url.strip():
        return None
    try:
        parsed = urlparse(url.strip())
        if parsed.scheme.lower() != 'rtsp':
            return None

        path = parsed.path.lstrip('/')
        if parsed.query:
            path += '?' + parsed.query

        return {
            'login': parsed.username or '',
            'pass': parsed.password or '',
            'ipaddress': parsed.hostname or '',
            'port': str(parsed.port) if parsed.port else '554',
            'path': path
        }
    except Exception as e:
        logger.warning(f"Не удалось распарсить RTSP URL {url}: {e}")
        return None


# ============================================================================
# CameraImportService
# ============================================================================

class CameraImportService:
    """Сервис импорта/экспорта камер."""

    def __init__(self, camera_service):
        self.camera_service = camera_service

    def import_from_excel(self, file_path: Path) -> Dict[str, Any]:
        """
        Импортирует камеры из Excel файла.

        Returns:
            {'success': bool, 'imported': int, 'skipped': int, 'errors': list}
        """
        try:
            import openpyxl
        except ImportError:
            return {'success': False, 'error': 'openpyxl не установлен. Выполните: pip install openpyxl'}

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active

            # Читаем заголовки
            headers = [str(cell.value).strip().lower() if cell.value else '' 
                      for cell in next(ws.iter_rows(min_row=1, max_row=1))]

            # Маппинг колонок
            col_map = {}
            found_fields = []
            for field, aliases in COLUMN_MAPPING.items():
                for alias in aliases:
                    alias_lower = alias.lower()
                    if alias_lower in headers:
                        col_map[field] = headers.index(alias_lower)
                        found_fields.append(field)
                        break

            # Проверяем обязательные поля
            required = [f for f, spec in CAMERA_FIELDS.items() if spec['required']]
            missing = [f for f in required if f not in col_map]
            if missing:
                return {
                    'success': False,
                    'error': f"Не найдены обязательные колонки: {', '.join(missing)}. Найдены: {', '.join(found_fields)}"
                }

            # Читаем данные
            cameras = []
            row_num = 1
            skipped_rows = 0
            errors = []

            for row in ws.iter_rows(min_row=2):
                row_num += 1
                values = [cell.value for cell in row]

                # Пропускаем полностью пустые строки
                if all(v is None or str(v).strip() == '' for v in values):
                    continue

                try:
                    camera = {}
                    for field, spec in CAMERA_FIELDS.items():
                        if field in col_map:
                            raw_value = values[col_map[field]]
                            if spec['type'] == 'str':
                                if field == 'name':
                                    camera[field] = clean_name(raw_value)
                                elif field in ('main_url', 'sub_url', 'sub2_url'):
                                    camera[field] = clean_url(raw_value)
                                else:
                                    camera[field] = clean_str(raw_value)
                            elif spec['type'] == 'bool':
                                camera[field] = parse_bool(raw_value, spec['default'])
                        else:
                            camera[field] = spec['default']

                    # Парсим полные RTSP URL (если login/pass/ip не указаны отдельно)
                    if camera.get('main_url') and camera['main_url'].strip().lower().startswith('rtsp://'):
                        parsed = parse_rtsp_url(camera['main_url'])
                        if parsed:
                            if not camera.get('login'):
                                camera['login'] = parsed['login']
                            if not camera.get('pass'):
                                camera['pass'] = parsed['pass']
                            if not camera.get('ipaddress'):
                                camera['ipaddress'] = parsed['ipaddress']
                            if not camera.get('port') or camera.get('port') == '554':
                                camera['port'] = parsed['port']
                            camera['main_url'] = parsed['path']

                    # Валидация обязательных полей
                    if not camera.get('id') or not camera.get('ipaddress') or not camera.get('main_url'):
                        errors.append(f"Строка {row_num}: пропущена (нет ID, IP или main_url)")
                        skipped_rows += 1
                        continue

                    # Если имя пустое — используем ID
                    if not camera['name']:
                        camera['name'] = camera['id']

                    # Если sub_url совпадает с main_url — оставляем пустым
                    if camera.get('sub_url') == camera.get('main_url'):
                        camera['sub_url'] = ''

                    cameras.append(camera)

                except Exception as e:
                    errors.append(f"Строка {row_num}: ошибка парсинга - {str(e)}")
                    skipped_rows += 1
                    continue

            wb.close()

            if not cameras:
                return {'success': False, 'error': 'Не найдено валидных камер в файле', 'errors': errors}

            # PATCH-193: МЕРЖ — обновление существующих + добавление новых
            current_cams = {c.id: c.to_dict() for c in self.camera_service.all_cameras()}
            updated = 0
            added = 0
            for new_cam in cameras:
                cam_id = new_cam.get('id')
                if cam_id in current_cams:
                    current_cams[cam_id].update(new_cam)
                    updated += 1
                else:
                    current_cams[cam_id] = new_cam
                    added += 1

            all_cams = list(current_cams.values())
            self.camera_service.save_cameras(all_cams)

            # PATCH-206: камеры вне наборов (включая отключённые) → в целевой набор
            linked, target = self._ensure_set_membership([c.get('id') for c in all_cams])

            return {
                'success': True,
                'imported': len(all_cams),
                'updated': updated,
                'added': added,
                'skipped': skipped_rows,
                'linked_to_set': linked,
                'target_set': target,
                'errors': errors
            }

        except Exception as e:
            return {'success': False, 'error': f'Ошибка импорта: {str(e)}'}

    def export_to_excel(self, file_path: Path) -> Dict[str, Any]:
        """Экспортирует камеры в Excel файл."""
        try:
            import openpyxl
        except ImportError:
            return {'success': False, 'error': 'openpyxl не установлен'}

        try:
            cameras = self.camera_service.all_cameras()

            # PATCH-193: проверка пустого списка
            if not cameras:
                logger.warning("export_to_excel: список камер пуст")
                # всё равно создаём файл с заголовками
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Камеры"
                headers = ['id', 'name', 'login', 'pass', 'ipaddress', 'port', 'main_url', 'sub_url', 'sub2_url', 'enabled', 'comment', 'audio', 'location']  # PATCH-197: имена колонок БД
                ws.append(headers)
                wb.save(file_path)
                wb.close()
                return {'success': True, 'exported': 0, 'warning': 'Список камер пуст'}

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Камеры"

            headers = ['id', 'name', 'login', 'pass', 'ipaddress', 'port', 'main_url', 'sub_url', 'sub2_url', 'enabled', 'comment', 'audio', 'location']  # PATCH-197: имена колонок БД
            ws.append(headers)

            for cam in cameras:
                ws.append([
                    cam.id, cam.name, cam.login, cam.pass_, cam.ipaddress, cam.port,
                    cam.main_url, cam.sub_url, cam.sub2_url,
                    'true' if cam.enabled else 'false',
                    cam.comment,
                    'true' if cam.audio else 'false',  # PATCH-198: текст, не locale-boolean
                    cam.location
                ])

            wb.save(file_path)
            wb.close()
            return {'success': True, 'exported': len(cameras)}

        except Exception as e:
            return {'success': False, 'error': f'Ошибка экспорта: {str(e)}'}

    def import_from_json(self, data: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Импортирует камеры из JSON массива (МЕРЖ: обновление + добавление)."""
        try:
            current_cams = {c.id: c.to_dict() for c in self.camera_service.all_cameras()}
            updated = 0
            added = 0
            errors = []
            warnings = []  # PATCH-226

            for i, cam_data in enumerate(data):
                try:
                    # Парсим полные RTSP URL
                    if cam_data.get('main_url') and cam_data['main_url'].strip().lower().startswith('rtsp://'):
                        parsed = parse_rtsp_url(cam_data['main_url'])
                        if parsed:
                            if not cam_data.get('login'):
                                cam_data['login'] = parsed['login']
                            if not cam_data.get('pass'):
                                cam_data['pass'] = parsed['pass']
                            if not cam_data.get('ipaddress'):
                                cam_data['ipaddress'] = parsed['ipaddress']
                            if not cam_data.get('port') or cam_data.get('port') == '554':
                                cam_data['port'] = parsed['port']
                            cam_data['main_url'] = parsed['path']

                    cam_id = cam_data.get('id')
                    if not cam_id or not cam_data.get('ipaddress') or not cam_data.get('main_url'):
                        errors.append(f"Камера {i}: пропущена (нет ID/IP/main_url)")
                        continue

                    # PATCH-226: валидация ID
                    id_warnings = _validate_camera_id(cam_id)
                    if id_warnings:
                        warnings.append({'id': cam_id, 'messages': id_warnings})

                    if cam_id in current_cams:
                        current_cams[cam_id].update(cam_data)
                        updated += 1
                    else:
                        current_cams[cam_id] = cam_data
                        added += 1

                except Exception as e:
                    errors.append(f"Камера {i}: ошибка - {str(e)}")
                    continue

            all_cams = list(current_cams.values())
            if not all_cams:
                return {'success': False, 'error': 'Не найдено валидных камер', 'errors': errors}

            self.camera_service.save_cameras(all_cams)

            # PATCH-206: камеры вне наборов (включая отключённые) → в целевой набор
            linked, target = self._ensure_set_membership([c.get('id') for c in all_cams])

            return {
                'success': True,
                'imported': len(all_cams),
                'updated': updated,
                'added': added,
                'linked_to_set': linked,
                'target_set': target,
                'errors': errors,
                'warnings': warnings,  # PATCH-226.3
            }

        except Exception as e:
            return {'success': False, 'error': f'Ошибка импорта: {str(e)}'}

    def _ensure_set_membership(self, camera_ids: List[str]):
        """PATCH-206: камеры, не состоящие ни в одном наборе (включая
        отключённые), допривязываются к целевому набору (самому наполненному).
        Возвращает (count, target_set_id)."""
        from app.db.repositories import set_repo
        sets = (set_repo.get_all() or {}).get("sets", {})
        if not sets:
            return 0, ""
        member = set()
        for s in sets.values():
            member.update(s.get("camera_ids", s.get("cameras", [])) or [])
        missing = [cid for cid in camera_ids if cid and cid not in member]
        if not missing:
            return 0, ""
        target = sorted(
            sets.keys(),
            key=lambda k: (-len(sets[k].get("camera_ids", sets[k].get("cameras", [])) or []), k)
        )[0]
        n = set_repo.add_cameras(target, missing)
        return n, target

    def export_to_json(self) -> List[Dict[str, Any]]:
        """Экспортирует камеры в JSON массив."""
        cameras = self.camera_service.all_cameras()
        return [cam.to_dict() for cam in cameras]


# ============================================================================
# Инициализация
# ============================================================================

from app.services.camera_service import camera_service

camera_import_service = CameraImportService(camera_service)
